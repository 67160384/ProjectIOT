"""Build the editable and printable project report from docs/REPORT.md."""
from __future__ import annotations

import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
SOURCE = DOCS / "REPORT.md"
DOCX_OUTPUT = DOCS / "รายงานโครงงาน-ฉบับพร้อมส่ง.docx"
HTML_OUTPUT = DOCS / "รายงานโครงงาน-ฉบับพร้อมส่ง.html"
PDF_OUTPUT = DOCS / "รายงานโครงงาน-ฉบับพร้อมส่ง.pdf"

GREEN = "164E43"
LIGHT_GREEN = "E7F2EE"
CREAM = "F7F3E8"
GRAY = "5D6965"


def set_cell_shading(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = props.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        props.append(shading)
    shading.set(qn("w:fill"), fill)


def keep_row(row) -> None:
    props = row._tr.get_or_add_trPr()
    props.append(OxmlElement("w:cantSplit"))


def repeat_header(row) -> None:
    props = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    props.append(element)


def add_page_number(paragraph) -> None:
    paragraph.add_run("หน้า ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph._p.extend((begin, instruction, separate, end))


def set_font(run, name="Noto Sans Thai", size=None, bold=None, color=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def configure(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2)

    style_specs = {
        "Normal": ("Noto Sans Thai", 10.5, False, "202926"),
        "Title": ("Noto Sans Thai", 23, True, GREEN),
        "Subtitle": ("Noto Sans Thai", 13, False, GRAY),
        "Heading 1": ("Noto Sans Thai", 16, True, GREEN),
        "Heading 2": ("Noto Sans Thai", 13, True, GREEN),
        "Heading 3": ("Noto Sans Thai", 11.5, True, GREEN),
        "Caption": ("Noto Sans Thai", 9, False, GRAY),
        "List Bullet": ("Noto Sans Thai", 10.5, False, "202926"),
        "List Number": ("Noto Sans Thai", 10.5, False, "202926"),
    }
    for style_name, (font, size, bold, color) in style_specs.items():
        style = document.styles[style_name]
        style.font.name = font
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
        style.element.get_or_add_rPr().rFonts.set(qn("w:cs"), font)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_after = Pt(5)
    document.styles["Normal"].paragraph_format.line_spacing = 1.18
    document.styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    document.styles["Heading 1"].paragraph_format.space_before = Pt(14)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(8)
    document.styles["Heading 1"].paragraph_format.keep_with_next = True
    document.styles["Heading 2"].paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "ใบเสร็จดี · Offline AI Receipt Analyzer"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_font(run, size=8.5, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)
    for run in footer.runs:
        set_font(run, size=8.5, color=GRAY)


def add_inline(paragraph, text: str) -> None:
    """Add a small Markdown subset: bold and inline code."""
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`)")
    for part in pattern.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, name="DejaVu Sans Mono", size=9, color=GREEN)
        else:
            run = paragraph.add_run(part)
            set_font(run)


def add_caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(paragraph, text)


def add_figure(document: Document, path: Path, caption: str, width_cm: float) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
    add_caption(document, caption)


def add_table(document: Document, lines: list[str]) -> None:
    parsed = [
        [cell.strip() for cell in line.strip().strip("|").split("|")]
        for line in lines
        if not re.fullmatch(r"\|?[\s:|\-]+\|?", line)
    ]
    if not parsed:
        return
    columns = max(len(row) for row in parsed)
    table = document.add_table(rows=0, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_index, data in enumerate(parsed):
        row = table.add_row()
        keep_row(row)
        if row_index == 0:
            repeat_header(row)
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = data[col_index] if col_index < len(data) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(1.5)
            add_inline(paragraph, text)
            for run in paragraph.runs:
                set_font(run, size=8.7, bold=(row_index == 0))
            if row_index == 0:
                set_cell_shading(cell, GREEN)
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F3F8F6")
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_note(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.5)
    paragraph.paragraph_format.right_indent = Cm(0.5)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(8)
    props = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), CREAM)
    props.append(shading)
    add_inline(paragraph, text)


def add_cover(document: Document) -> None:
    for _ in range(4):
        document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("รายงานโครงงาน")
    set_font(run, size=17, bold=True, color=GREEN)
    project = document.add_paragraph()
    project.alignment = WD_ALIGN_PARAGRAPH.CENTER
    project.paragraph_format.space_after = Pt(8)
    run = project.add_run("ใบเสร็จดี")
    set_font(run, size=28, bold=True, color=GREEN)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(subtitle, "Offline AI Receipt Analyzer\nระบบวิเคราะห์ใบเสร็จและสรุปค่าใช้จ่ายแบบออฟไลน์บน Raspberry Pi")
    for run in subtitle.runs:
        set_font(run, size=13, color=GRAY)

    document.add_paragraph()
    add_figure(document, DOCS / "architecture.png", "ภาพรวมการทำงานของโครงงาน", 13.8)
    for _ in range(2):
        document.add_paragraph()
    for text in (
        "ผู้จัดทำ  [กรอกชื่อ–นามสกุล]   รหัสนักศึกษา [กรอก]",
        "รายวิชา [กรอก]   อาจารย์ผู้สอน [กรอก]",
        "สาขาวิชา/สถานศึกษา [กรอก]   ภาคเรียน [กรอก]",
        "จัดทำเมื่อ 25 ตุลาคม 2569",
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(paragraph, text)
    document.add_page_break()


def build_docx() -> None:
    document = Document()
    configure(document)
    add_cover(document)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    # The cover already contains the title and metadata. Start at the abstract.
    start = next(index for index, line in enumerate(lines) if line == "## บทคัดย่อ")
    lines = lines[start:]
    table_lines: list[str] = []
    code_lines: list[str] = []
    in_code = False
    figure_number = 2

    def flush_table() -> None:
        nonlocal table_lines
        if table_lines:
            add_table(document, table_lines)
            table_lines = []

    def flush_code() -> None:
        nonlocal code_lines
        if not code_lines:
            return
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.45)
        paragraph.paragraph_format.right_indent = Cm(0.45)
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(7)
        props = paragraph._p.get_or_add_pPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F0F3F2")
        props.append(shading)
        run = paragraph.add_run("\n".join(code_lines))
        set_font(run, name="DejaVu Sans Mono", size=8.2, color="19332E")
        code_lines = []

    for line in lines:
        if line.startswith("```"):
            flush_table()
            if in_code:
                flush_code()
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
            continue
        if line.startswith("|"):
            table_lines.append(line)
            continue
        flush_table()
        stripped = line.strip()
        if not stripped:
            continue
        if stripped == "ภาพประกอบหลัก: architecture.png":
            add_figure(document, DOCS / "architecture.png", f"ภาพที่ {figure_number} สถาปัตยกรรมและลำดับการทำงานของระบบ", 16.1)
            figure_number += 1
            continue
        screenshot = re.fullmatch(r"ภาพหน้าจอระบบ: (.+\.png)", stripped)
        if screenshot:
            file_name = screenshot.group(1)
            captions = {
                "01-upload.png": "หน้าอัปโหลดภาพใบเสร็จ",
                "08-ai-before-review.png": "ข้อมูลที่ระบบกรอกอัตโนมัติก่อนผู้ใช้ตรวจแก้",
                "03-dashboard.png": "Dashboard สรุปรายจ่ายและประวัติ",
            }
            width = 15.8 if file_name != "08-ai-before-review.png" else 13.1
            add_figure(document, DOCS / "evidence" / file_name, f"ภาพที่ {figure_number} {captions[file_name]} (ข้อมูลสาธิต)", width)
            figure_number += 1
            continue
        if stripped.startswith("## "):
            heading = document.add_heading(stripped[3:], level=1)
            if stripped != "## บทคัดย่อ":
                heading.paragraph_format.page_break_before = True
            continue
        if stripped.startswith("### "):
            document.add_heading(stripped[4:], level=2)
            continue
        if stripped.startswith("> "):
            add_note(document, stripped[2:])
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, numbered.group(1))
            continue
        if stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:])
            continue
        paragraph = document.add_paragraph()
        add_inline(paragraph, stripped)

    flush_table()
    flush_code()
    document.core_properties.title = "รายงานโครงงาน ใบเสร็จดี — Offline AI Receipt Analyzer"
    document.core_properties.subject = "ระบบวิเคราะห์ใบเสร็จแบบออฟไลน์บน Raspberry Pi"
    document.core_properties.author = "[กรอกชื่อผู้จัดทำ]"
    document.core_properties.keywords = "Raspberry Pi, OCR, Local LLM, Flask, SQLite"
    document.save(DOCX_OUTPUT)


def build_html() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    replacements = {
        "ภาพประกอบหลัก: architecture.png": "![สถาปัตยกรรมและลำดับการทำงานของระบบ](architecture.png){ width=92% }",
        "ภาพหน้าจอระบบ: 01-upload.png": "![หน้าอัปโหลดภาพใบเสร็จ — ข้อมูลสาธิต](evidence/01-upload.png){ width=92% }",
        "ภาพหน้าจอระบบ: 08-ai-before-review.png": "![ข้อมูลที่ระบบกรอกอัตโนมัติก่อนผู้ใช้ตรวจแก้ — ข้อมูลสาธิต](evidence/08-ai-before-review.png){ width=74% }",
        "ภาพหน้าจอระบบ: 03-dashboard.png": "![Dashboard สรุปรายจ่ายและประวัติ — ข้อมูลสาธิต](evidence/03-dashboard.png){ width=92% }",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    css_text = """
@font-face { font-family: NotoThai; src: url('file:///usr/share/fonts/truetype/noto/NotoSansThai-Regular.ttf'); }
@font-face { font-family: NotoThai; src: url('file:///usr/share/fonts/truetype/noto/NotoSansThai-Bold.ttf'); font-weight: 700; }
@page { size: A4; margin: 18mm 18mm 18mm 20mm; }
body { font-family: NotoThai, sans-serif; color: #202926; font-size: 10.5pt; line-height: 1.45; }
h1 { color:#164e43; font-size:25pt; text-align:center; margin-top:58mm; page-break-after:avoid; }
h1 + p { text-align:center; font-size:13pt; color:#5d6965; }
h2 { color:#164e43; font-size:16pt; margin-top:0; padding-top:3mm; page-break-before:always; page-break-after:avoid; border-bottom:1px solid #b9d3ca; }
h2:first-of-type { page-break-before:always; }
h3 { color:#164e43; font-size:12.5pt; page-break-after:avoid; }
p { text-align:justify; orphans:3; widows:3; }
li { margin-bottom:1.5mm; }
table { border-collapse:collapse; width:100%; margin:4mm 0; font-size:8.7pt; page-break-inside:auto; }
tr { page-break-inside:avoid; }
th { background:#164e43; color:white; font-weight:700; }
th, td { border:1px solid #9bb4ab; padding:2mm; vertical-align:top; }
tr:nth-child(even) td { background:#f3f8f6; }
pre { background:#f0f3f2; border-left:3px solid #4d8b79; padding:3mm; white-space:pre-wrap; font-size:8.4pt; page-break-inside:avoid; }
code { color:#164e43; }
blockquote { background:#f7f3e8; border-left:3px solid #c49a45; padding:3mm 4mm; margin:4mm 0; }
img { display:block; max-height:210mm; margin:4mm auto 1mm; page-break-inside:avoid; }
figure { page-break-inside:avoid; text-align:center; }
figcaption { font-size:8.7pt; color:#5d6965; }
strong { color:#173d36; }
""".strip()
    with tempfile.TemporaryDirectory(prefix="projectiot-report-") as temp_dir:
        prepared = Path(temp_dir) / "report-print.md"
        css = Path(temp_dir) / "report-print.css"
        prepared.write_text(text, encoding="utf-8")
        css.write_text(css_text, encoding="utf-8")
        subprocess.run(
            [
                "pandoc", str(prepared), "--standalone", "--embed-resources",
                "--resource-path", str(DOCS), "--css", str(css), "-o", str(HTML_OUTPUT),
                "--metadata", "lang=th", "--metadata", "title=รายงานโครงงาน ใบเสร็จดี",
            ],
            check=True,
            cwd=DOCS,
        )


def main() -> None:
    build_docx()
    build_html()
    print(f"Created {DOCX_OUTPUT.name} and {HTML_OUTPUT.name}")
    print("Print the HTML with Chromium to create the PDF.")


if __name__ == "__main__":
    main()
