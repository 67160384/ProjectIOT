"""Generate a clear two-page A4 user manual from real application screenshots."""
from __future__ import annotations

import base64
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
EVIDENCE = DOCS / "evidence"
DOCX_OUTPUT = DOCS / "คู่มือใช้งาน.docx"
HTML_OUTPUT = DOCS / "คู่มือใช้งาน.html"

GREEN = "0A4A37"
LIGHT = "EDF5EF"
CREAM = "F8F5EC"
GRAY = "5C6864"


def set_font(run, size=10, bold=False, color="202926"):
    name = "Noto Sans Thai"
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    props.append(element)


def add_text(paragraph, text, *, size=10, bold=False, color="202926"):
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return run


def add_step(document, number, title, detail):
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(1.2)
    table.columns[1].width = Cm(16.2)
    number_cell, text_cell = table.rows[0].cells
    number_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(number_cell, GREEN)
    p = number_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, str(number), size=12, bold=True, color="FFFFFF")
    p = text_cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    add_text(p, title, size=10.5, bold=True, color=GREEN)
    p = text_cell.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    add_text(p, detail, size=9.2, color=GRAY)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_heading(document, kicker, title, subtitle):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    add_text(p, kicker.upper(), size=8.5, bold=True, color=GRAY)
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    add_text(p, title, size=20, bold=True, color=GREEN)
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    add_text(p, subtitle, size=10, color=GRAY)


def add_caption(document, text):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    add_text(p, text, size=8.5, color=GRAY)


def configure_docx(document):
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.15)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "ใบเสร็จดี · คู่มือใช้งานระบบ Offline AI Receipt Analyzer", size=8, color=GRAY)


def build_docx():
    document = Document()
    configure_docx(document)

    add_page_heading(
        document,
        "คู่มือใช้งาน · หน้า 1/2",
        "เริ่มต้นใช้งานและอัปโหลดใบเสร็จ",
        "ใช้งานจาก Raspberry Pi หรืออุปกรณ์ที่เชื่อมต่อ Wi-Fi เดียวกัน",
    )
    add_step(document, 1, "เปิดระบบ", "เปิด Local AI ด้วย scripts/start_llm.sh และเปิดเว็บไซต์ด้วย scripts/start_web.sh จากโฟลเดอร์ ProjectIOT")
    add_step(document, 2, "เข้าเว็บไซต์", "บน Pi เปิด http://127.0.0.1:5000 หรือจากเครื่องอื่นเปิด http://<IP ของ Pi>:5000")
    add_step(document, 3, "เลือกภาพใบเสร็จ", "กด “เลือกรูปใบเสร็จของคุณ” เลือกไฟล์ JPG/PNG ที่เห็นรายการและยอดรวมครบ แล้วกด “เริ่มอ่านใบเสร็จ”")

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.add_run().add_picture(str(EVIDENCE / "01-upload.png"), width=Cm(16.9))
    add_caption(document, "ภาพที่ 1 หน้าอัปโหลด — เลือกภาพที่ตรง ไม่เบลอ และไม่มีแสงสะท้อน")

    note = document.add_table(rows=1, cols=1)
    note.alignment = WD_TABLE_ALIGNMENT.CENTER
    shade(note.cell(0, 0), CREAM)
    p = note.cell(0, 0).paragraphs[0]
    add_text(p, "ระหว่างประมวลผล", size=9.5, bold=True, color=GREEN)
    add_text(p, " ระบบจะปรับภาพ อ่าน OCR และวิเคราะห์ด้วย Local AI บน Raspberry Pi อาจใช้เวลาหลายวินาทีถึงหลายนาที ห้ามปิด terminal ของ AI และเว็บ", size=9.2)

    document.add_page_break()
    add_page_heading(
        document,
        "คู่มือใช้งาน · หน้า 2/2",
        "ตรวจแก้ ยืนยัน และดู Dashboard",
        "ข้อมูลจาก AI เป็นร่าง ผู้ใช้ต้องเทียบกับภาพต้นฉบับก่อนบันทึกทุกครั้ง",
    )

    layout = document.add_table(rows=1, cols=2)
    layout.alignment = WD_TABLE_ALIGNMENT.CENTER
    layout.autofit = False
    layout.columns[0].width = Cm(8.7)
    layout.columns[1].width = Cm(9.0)
    left, right = layout.rows[0].cells
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(EVIDENCE / "02-review.png"), width=Cm(7.7))
    p = left.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "ภาพที่ 2 หน้าตรวจแก้ข้อมูล", size=8.5, color=GRAY)

    for number, title, detail in (
        (4, "ตรวจข้อมูล", "ตรวจชื่อร้าน/ลูกค้า วันที่ ชื่อสินค้า จำนวน และยอดรวมบรรทัด แก้ไขช่องที่ผิดได้ทันที"),
        (5, "กรอกหมวดหมู่และส่วนลด", "พิมพ์หมวดหมู่เอง หากลด 6 บาทให้กรอก 6.00 ระบบจะนำไปลบจากยอดสินค้า"),
        (6, "ยืนยันและบันทึก", "ตรวจยอดสุทธิให้ตรงใบเสร็จ แล้วกด “ยืนยันและบันทึกค่าใช้จ่าย”"),
    ):
        box = right.add_table(rows=1, cols=2)
        box.autofit = False
        box.columns[0].width = Cm(1.0)
        box.columns[1].width = Cm(7.6)
        shade(box.cell(0, 0), GREEN)
        p = box.cell(0, 0).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, str(number), size=11, bold=True, color="FFFFFF")
        p = box.cell(0, 1).paragraphs[0]
        add_text(p, title, size=10, bold=True, color=GREEN)
        p = box.cell(0, 1).add_paragraph()
        add_text(p, detail, size=8.8, color=GRAY)
        right.add_paragraph().paragraph_format.space_after = Pt(1)

    p = right.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(EVIDENCE / "03-dashboard.png"), width=Cm(8.1))
    p = right.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "ภาพที่ 3 Dashboard แสดงยอดรวม หมวดหมู่ และประวัติ", size=8.3, color=GRAY)

    note = document.add_table(rows=1, cols=1)
    shade(note.cell(0, 0), LIGHT)
    p = note.cell(0, 0).paragraphs[0]
    add_text(p, "ข้อควรจำ: ", size=9.2, bold=True, color=GREEN)
    add_text(p, "ยอดบรรทัดคือยอดรวมของสินค้านั้นแล้ว ไม่ต้องคูณจำนวนซ้ำ • หาก AI อ่านผิดให้แก้เองก่อนบันทึก • Dashboard นับเฉพาะใบเสร็จที่ยืนยันแล้ว", size=9.0)

    document.core_properties.title = "คู่มือใช้งาน ใบเสร็จดี"
    document.core_properties.subject = "คู่มือระบบ Offline AI Receipt Analyzer"
    document.save(DOCX_OUTPUT)


def image_uri(path):
    return "data:image/png;base64," + base64.b64encode(path.read_bytes()).decode("ascii")


def build_html():
    upload = image_uri(EVIDENCE / "01-upload.png")
    review = image_uri(EVIDENCE / "02-review.png")
    dashboard = image_uri(EVIDENCE / "03-dashboard.png")
    html = f"""<!doctype html><html lang="th"><head><meta charset="utf-8"><title>คู่มือใช้งาน ใบเสร็จดี</title>
<style>
@font-face{{font-family:NotoThai;src:url('file:///usr/share/fonts/truetype/noto/NotoSansThai-Regular.ttf')}}
@font-face{{font-family:NotoThai;src:url('file:///usr/share/fonts/truetype/noto/NotoSansThai-Bold.ttf');font-weight:700}}
@page{{size:A4;margin:11mm 13mm}}*{{box-sizing:border-box}}body{{margin:0;font:10pt/1.42 NotoThai,sans-serif;color:#202926}}
.page{{height:275mm;page-break-after:always;position:relative}}.page:last-child{{page-break-after:auto}}.kicker{{font-size:8pt;letter-spacing:.08em;color:#5c6864;font-weight:700}}
h1{{font-size:21pt;color:#0a4a37;margin:1mm 0}}.sub{{color:#5c6864;margin:0 0 4mm}}.steps{{display:grid;grid-template-columns:repeat(3,1fr);gap:3mm;margin:3mm 0}}
.step{{border:1px solid #cbdad4;border-radius:3mm;padding:3mm;background:#fff}}.num{{display:inline-grid;place-items:center;width:8mm;height:8mm;border-radius:50%;background:#0a4a37;color:#fff;font-weight:700;margin-right:2mm}}
.step b{{color:#0a4a37}}.step p{{font-size:8.5pt;margin:2mm 0 0}}.shot{{width:100%;max-height:154mm;object-fit:contain;border:1px solid #d7e0dc;border-radius:3mm;background:#f5f6f2}}
.caption{{font-size:8.3pt;text-align:center;color:#5c6864;margin:1mm 0 3mm}}.note{{background:#f8f5ec;border-left:3px solid #c69a47;padding:3mm 4mm;border-radius:2mm;font-size:8.8pt}}
.two{{display:grid;grid-template-columns:48% 52%;gap:5mm;align-items:start}}.review{{width:100%;max-height:205mm;object-fit:contain;border:1px solid #d7e0dc;border-radius:3mm}}
.right .step{{margin-bottom:3mm}}.dash{{width:100%;border:1px solid #d7e0dc;border-radius:3mm;margin-top:1mm}}.remember{{background:#edf5ef;padding:3mm;border-radius:2mm;font-size:8.7pt;margin-top:3mm}}
footer{{position:absolute;bottom:0;left:0;right:0;text-align:center;font-size:7.5pt;color:#7b8581}}
</style></head><body>
<section class="page"><div class="kicker">คู่มือใช้งาน · หน้า 1/2</div><h1>เริ่มต้นใช้งานและอัปโหลดใบเสร็จ</h1><p class="sub">ใช้งานจาก Raspberry Pi หรืออุปกรณ์ที่เชื่อมต่อ Wi-Fi เดียวกัน</p>
<div class="steps"><div class="step"><span class="num">1</span><b>เปิดระบบ</b><p>เปิด Local AI ด้วย <b>start_llm.sh</b> และเว็บไซต์ด้วย <b>start_web.sh</b></p></div>
<div class="step"><span class="num">2</span><b>เข้าเว็บไซต์</b><p>บน Pi เปิด 127.0.0.1:5000 หรือใช้ IP ของ Pi จากเครื่องอื่น</p></div>
<div class="step"><span class="num">3</span><b>เลือกใบเสร็จ</b><p>เลือก JPG/PNG ที่เห็นรายการกับยอดครบ แล้วกด “เริ่มอ่านใบเสร็จ”</p></div></div>
<img class="shot" src="{upload}"><div class="caption">ภาพที่ 1 หน้าอัปโหลด — ภาพควรตรง คมชัด และไม่มีแสงสะท้อน</div>
<div class="note"><b>ระหว่างประมวลผล:</b> ระบบจะปรับภาพ อ่าน OCR และวิเคราะห์ด้วย Local AI บน Raspberry Pi อาจใช้เวลาหลายวินาทีถึงหลายนาที ห้ามปิด terminal ของ AI และเว็บ</div>
<footer>ใบเสร็จดี · Offline AI Receipt Analyzer</footer></section>
<section class="page"><div class="kicker">คู่มือใช้งาน · หน้า 2/2</div><h1>ตรวจแก้ ยืนยัน และดู Dashboard</h1><p class="sub">ข้อมูลจาก AI เป็นร่าง ผู้ใช้ต้องเทียบกับภาพต้นฉบับก่อนบันทึกทุกครั้ง</p>
<div class="two"><div><img class="review" src="{review}"><div class="caption">ภาพที่ 2 หน้าตรวจแก้ข้อมูล</div></div><div class="right">
<div class="step"><span class="num">4</span><b>ตรวจข้อมูล</b><p>ตรวจชื่อร้าน/ลูกค้า วันที่ สินค้า จำนวน และยอดรวมบรรทัด แก้ไขช่องที่ผิดได้ทันที</p></div>
<div class="step"><span class="num">5</span><b>กรอกหมวดหมู่และส่วนลด</b><p>พิมพ์หมวดหมู่เอง หากลด 6 บาทให้กรอก 6.00 ระบบจะนำไปลบจากยอดสินค้า</p></div>
<div class="step"><span class="num">6</span><b>ยืนยันและบันทึก</b><p>ตรวจยอดสุทธิให้ตรงใบเสร็จ แล้วกด “ยืนยันและบันทึกค่าใช้จ่าย”</p></div>
<img class="dash" src="{dashboard}"><div class="caption">ภาพที่ 3 Dashboard แสดงยอดรวม หมวดหมู่ และประวัติ</div>
<div class="remember"><b>ข้อควรจำ:</b> ยอดบรรทัดเป็นยอดรวมแล้ว ไม่ต้องคูณจำนวนซ้ำ • หาก AI อ่านผิดให้แก้ก่อนบันทึก • Dashboard นับเฉพาะใบที่ยืนยันแล้ว</div>
</div></div><footer>ใบเสร็จดี · Offline AI Receipt Analyzer</footer></section></body></html>"""
    HTML_OUTPUT.write_text(html, encoding="utf-8")


if __name__ == "__main__":
    build_docx()
    build_html()
    print(f"Created {DOCX_OUTPUT.name} and {HTML_OUTPUT.name}")
