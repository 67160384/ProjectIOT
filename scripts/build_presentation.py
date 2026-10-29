"""Build a concise Thai presentation for the Offline AI Receipt Analyzer."""
from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
EVIDENCE = DOCS / "evidence"
OUTPUT = DOCS / "สไลด์นำเสนอ-ใบเสร็จดี.pptx"

W = 13.333
H = 7.5
GREEN = "0A4A37"
GREEN_2 = "15664B"
PALE = "EAF2ED"
CREAM = "F7F4EA"
INK = "17231F"
GRAY = "60706A"
LINE = "CAD9D3"
LIME = "94B853"
GOLD = "D5A84B"
WHITE = "FFFFFF"
RED = "A75643"
FONT = "Noto Sans Thai"
LATIN = "Noto Sans"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_run(run, size=18, bold=False, color=INK, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)


def rect(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(1)
    return shape


def text(slide, value, x, y, w, h, size=18, bold=False, color=INK,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.03, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def rich_text(slide, parts, x, y, w, h, size=18, align=PP_ALIGN.LEFT,
              valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.03)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    for content, bold, color in parts:
        run = p.add_run()
        run.text = content
        set_run(run, size=size, bold=bold, color=color)
    return box


def bullet_list(slide, items, x, y, w, h, size=16, color=INK, accent=GREEN):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.02)
    for index, item in enumerate(items):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)
        p.text = f"●  {item}"
        for run in p.runs:
            set_run(run, size=size, color=color)
        if p.runs:
            p.runs[0].font.color.rgb = rgb(accent)
    return box


def add_picture_contain(slide, path, x, y, w, h):
    with Image.open(path) as image:
        iw, ih = image.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    return slide.shapes.add_picture(
        str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
        width=Inches(pw), height=Inches(ph),
    )


def add_picture_cover(slide, path, x, y, w, h):
    with Image.open(path) as image:
        iw, ih = image.size
    image_ratio = iw / ih
    box_ratio = w / h
    if image_ratio > box_ratio:
        ph = h
        pw = h * image_ratio
    else:
        pw = w
        ph = w / image_ratio
    picture = slide.shapes.add_picture(
        str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
        width=Inches(pw), height=Inches(ph),
    )
    return picture


def circle_number(slide, number, x, y, diameter=0.52, fill=GREEN):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(diameter), Inches(diameter))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background()
    text(slide, str(number), x, y + 0.01, diameter, diameter - 0.02, size=14, bold=True,
         color=WHITE, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    return shape


def header(slide, kicker, title, page):
    text(slide, kicker.upper(), 0.7, 0.34, 8.8, 0.25, size=8.5, bold=True, color=GRAY)
    text(slide, title, 0.7, 0.66, 11.4, 0.55, size=25, bold=True, color=GREEN)
    text(slide, f"{page:02d}", 12.1, 0.46, 0.55, 0.35, size=10, bold=True, color=GRAY, align=PP_ALIGN.RIGHT)
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.7), Inches(1.28), Inches(12.65), Inches(1.28))
    line.line.color.rgb = rgb(LINE)
    line.line.width = Pt(1)


def footer(slide):
    text(slide, "ใบเสร็จดี · Offline AI Receipt Analyzer", 0.7, 7.13, 5.0, 0.18, size=7.5, color=GRAY)


def blank_slide(prs, background=WHITE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(background)
    return slide


def slide_cover(prs):
    slide = blank_slide(prs, CREAM)
    rect(slide, 7.15, 0.55, 5.55, 6.4, fill=WHITE, line=LINE)
    add_picture_contain(slide, DOCS / "dashboard-preview.png", 7.35, 0.75, 5.15, 6.0)
    text(slide, "PROJECT IOT · EDGE AI", 0.75, 0.75, 5.8, 0.3, size=10, bold=True, color=GREEN_2)
    text(slide, "ใบเสร็จดี", 0.75, 1.42, 5.9, 0.72, size=34, bold=True, color=GREEN)
    text(slide, "Offline AI Receipt Analyzer", 0.75, 2.16, 5.9, 0.42, size=19, bold=True, color=INK)
    text(slide, "ระบบอ่านใบเสร็จ สรุปค่าใช้จ่าย และแสดง Dashboard\nด้วย AI ที่ทำงานภายใน Raspberry Pi", 0.75, 2.86, 5.95, 1.05, size=16, color=GRAY)
    labels = [("Raspberry Pi 4", GREEN), ("OCR ไทย–อังกฤษ", GREEN_2), ("Local LLM", LIME)]
    x = 0.75
    for label, color in labels:
        shape = rect(slide, x, 4.28, 1.76, 0.42, fill=WHITE, line=color)
        text(slide, label, x, 4.34, 1.76, 0.26, size=9.5, bold=True, color=color, align=PP_ALIGN.CENTER)
        x += 1.9
    text(slide, "ผู้จัดทำ: [ชื่อ–นามสกุล]  |  รหัสนักศึกษา: [กรอก]", 0.75, 5.65, 6.1, 0.35, size=12, color=INK)
    text(slide, "นำเสนอเดือนตุลาคม 2569", 0.75, 6.1, 4.0, 0.3, size=10, color=GRAY)


def slide_problem(prs):
    slide = blank_slide(prs)
    header(slide, "01 · Problem & Goal", "ปัญหาและเป้าหมายของโครงงาน", 2)
    cards = [
        ("บันทึกด้วยมือ", "ต้องคัดลอกชื่อสินค้าและราคา\nใช้เวลาและพิมพ์ผิดได้", GOLD),
        ("ข้อมูลส่วนตัว", "ใบเสร็จอาจมีชื่อ ที่อยู่\nและประวัติการซื้อ", RED),
        ("รูปแบบหลากหลาย", "ตาราง คำย่อ และส่วนลด\nอยู่คนละตำแหน่ง", GREEN_2),
    ]
    for index, (title_value, detail, color) in enumerate(cards):
        x = 0.75 + index * 4.12
        rect(slide, x, 1.65, 3.72, 1.7, fill=CREAM if index == 0 else PALE, line=color)
        circle_number(slide, index + 1, x + 0.23, 1.88, fill=color)
        text(slide, title_value, x + 0.85, 1.84, 2.55, 0.35, size=17, bold=True, color=INK)
        text(slide, detail, x + 0.25, 2.38, 3.2, 0.7, size=12.5, color=GRAY)
    rect(slide, 0.75, 3.78, 11.85, 2.48, fill=GREEN, line=GREEN)
    text(slide, "เป้าหมาย", 1.05, 4.05, 2.0, 0.42, size=20, bold=True, color=WHITE)
    goals = [
        "อ่านข้อความไทยและอังกฤษจากภาพใบเสร็จ",
        "กรอกข้อมูลอัตโนมัติ แต่ให้ผู้ใช้ตรวจแก้ก่อนบันทึก",
        "เก็บข้อมูลและแสดง Dashboard โดยไม่ส่งภาพไป Cloud AI",
    ]
    for index, goal in enumerate(goals):
        circle_number(slide, "✓", 3.25, 4.0 + index * 0.62, diameter=0.42, fill=LIME)
        text(slide, goal, 3.82, 4.01 + index * 0.62, 7.95, 0.4, size=14.5, color=WHITE)
    footer(slide)


def slide_architecture(prs):
    slide = blank_slide(prs)
    header(slide, "02 · Architecture", "ภาพรวมโครงสร้างระบบ", 3)
    rect(slide, 0.7, 1.48, 11.95, 4.95, fill=WHITE, line=LINE)
    add_picture_contain(slide, DOCS / "architecture.png", 0.95, 1.7, 11.45, 4.35)
    rect(slide, 3.55, 6.1, 6.25, 0.62, fill=PALE, line=GREEN_2)
    text(slide, "โทรศัพท์เชื่อมผ่าน LAN · OCR, LLM, ฐานข้อมูล และเว็บอยู่บน Raspberry Pi", 3.7, 6.25, 5.95, 0.24, size=11.5, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    footer(slide)


def slide_workflow(prs):
    slide = blank_slide(prs)
    header(slide, "03 · Workflow", "การทำงานตั้งแต่รูปภาพถึง Dashboard", 4)
    steps = [
        ("อัปโหลด", "JPG / PNG"),
        ("ปรับภาพ", "หมุน · ขยาย · เพิ่มความคม"),
        ("OCR", "อ่านข้อความไทย–อังกฤษ"),
        ("Parser + LLM", "แยกร้าน สินค้า จำนวน ราคา"),
        ("ตรวจแก้", "ผู้ใช้ยืนยันข้อมูล"),
        ("บันทึก", "SQLite → Dashboard"),
    ]
    for index, (title_value, detail) in enumerate(steps):
        row = index // 3
        col = index % 3
        x = 0.78 + col * 4.14
        y = 1.68 + row * 2.25
        rect(slide, x, y, 3.65, 1.55, fill=PALE if row == 0 else CREAM, line=GREEN_2)
        circle_number(slide, index + 1, x + 0.2, y + 0.22, fill=GREEN if index != 3 else LIME)
        text(slide, title_value, x + 0.86, y + 0.2, 2.45, 0.38, size=18, bold=True, color=GREEN)
        text(slide, detail, x + 0.22, y + 0.82, 3.15, 0.46, size=12.2, color=GRAY, align=PP_ALIGN.CENTER)
        if col < 2:
            connector = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(x + 3.65), Inches(y + 0.78), Inches(x + 4.05), Inches(y + 0.78),
            )
            connector.line.color.rgb = rgb(GREEN_2)
            connector.line.width = Pt(2)
    text(slide, "AI สร้าง “ร่างข้อมูล” เท่านั้น — ผู้ใช้เป็นผู้ยืนยันผลสุดท้าย", 2.45, 6.38, 8.45, 0.38, size=15, bold=True, color=RED, align=PP_ALIGN.CENTER)
    footer(slide)


def slide_technology(prs):
    slide = blank_slide(prs)
    header(slide, "04 · Technology", "เครื่องมือที่ใช้และหน้าที่", 5)
    groups = [
        ("อุปกรณ์", [("Raspberry Pi 4", "ประมวลผลและเก็บข้อมูล"), ("Wi-Fi / LAN", "เชื่อมโทรศัพท์กับ Pi")], GREEN),
        ("Backend", [("Flask + Waitress", "เว็บและควบคุม workflow"), ("SQLite", "ฐานข้อมูลใบเสร็จ")], GREEN_2),
        ("OCR & AI", [("Pillow + Tesseract", "ปรับภาพและอ่านข้อความ"), ("Qwen + llama.cpp", "วิเคราะห์ข้อความในเครื่อง")], LIME),
        ("หน้าจอ", [("HTML / Bootstrap", "ส่วนติดต่อผู้ใช้"), ("JavaScript / Chart.js", "คำนวณและแสดงกราฟ")], GOLD),
    ]
    for index, (group, rows, color) in enumerate(groups):
        col = index % 2
        row = index // 2
        x = 0.75 + col * 6.05
        y = 1.55 + row * 2.48
        rect(slide, x, y, 5.65, 2.08, fill=WHITE, line=color)
        rect(slide, x, y, 1.45, 2.08, fill=color, line=color)
        text(slide, group, x + 0.08, y + 0.7, 1.28, 0.5, size=14.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        for r_index, (name, role) in enumerate(rows):
            yy = y + 0.3 + r_index * 0.88
            text(slide, name, x + 1.68, yy, 2.05, 0.48, size=11.8, bold=True, color=INK, font=LATIN)
            text(slide, role, x + 3.72, yy, 1.55, 0.52, size=10.7, color=GRAY)
    footer(slide)


def slide_llm(prs):
    slide = blank_slide(prs)
    header(slide, "05 · Local LLM", "LLM เกี่ยวข้องกับระบบอย่างไร", 6)
    stages = [
        ("Tesseract OCR", "อ่านตัวอักษร\nจากภาพ", GREEN_2),
        ("Python Parser", "จับวันที่ ราคา\nและรูปแบบตาราง", GOLD),
        ("Qwen2.5-0.5B", "เข้าใจว่าข้อความ\nคือข้อมูลประเภทใด", LIME),
        ("JSON → ฟอร์ม", "กรอกข้อมูลให้\nผู้ใช้ตรวจแก้", GREEN),
    ]
    for index, (name, role, color) in enumerate(stages):
        x = 0.72 + index * 3.12
        rect(slide, x, 1.72, 2.62, 1.85, fill=PALE if index % 2 == 0 else CREAM, line=color)
        text(slide, name, x + 0.12, 2.0, 2.38, 0.38, size=16, bold=True, color=color, align=PP_ALIGN.CENTER)
        text(slide, role, x + 0.18, 2.52, 2.26, 0.68, size=12, color=GRAY, align=PP_ALIGN.CENTER)
        if index < 3:
            text(slide, "→", x + 2.65, 2.35, 0.42, 0.4, size=23, bold=True, color=GREEN_2, align=PP_ALIGN.CENTER)
    rect(slide, 0.75, 4.05, 12.0, 1.9, fill=GREEN, line=GREEN)
    text(slide, "ประเด็นสำคัญ", 1.08, 4.34, 2.2, 0.4, size=19, bold=True, color=WHITE)
    bullet_list(slide, [
        "LLM รับ “ข้อความจาก OCR” ไม่ได้รับรูปภาพโดยตรง",
        "รันผ่าน llama.cpp ภายใน Raspberry Pi",
        "ใช้โมเดลแบบ quantized เพื่อให้เหมาะกับ RAM 4 GB",
    ], 3.25, 4.25, 8.75, 1.35, size=13.5, color=WHITE, accent=LIME)
    footer(slide)


def slide_screens(prs):
    slide = blank_slide(prs)
    header(slide, "06 · User Interface", "หน้าจอหลักของระบบ", 7)
    cards = [
        ("1 · อัปโหลด", EVIDENCE / "01-upload.png", "เลือกภาพใบเสร็จ"),
        ("2 · ตรวจแก้", EVIDENCE / "02-review.png", "แก้ข้อมูลก่อนบันทึก"),
        ("3 · Dashboard", EVIDENCE / "03-dashboard.png", "ดูยอดและประวัติ"),
    ]
    for index, (title_value, image, caption) in enumerate(cards):
        x = 0.68 + index * 4.18
        rect(slide, x, 1.55, 3.82, 4.88, fill=WHITE, line=LINE)
        text(slide, title_value, x + 0.2, 1.77, 3.42, 0.35, size=16, bold=True, color=GREEN)
        add_picture_contain(slide, image, x + 0.18, 2.2, 3.46, 3.55)
        text(slide, caption, x + 0.22, 5.86, 3.38, 0.3, size=10.8, color=GRAY, align=PP_ALIGN.CENTER)
    footer(slide)


def slide_results(prs):
    slide = blank_slide(prs)
    header(slide, "07 · Results", "ผลที่ได้และข้อจำกัด", 8)
    rect(slide, 0.75, 1.58, 5.68, 4.95, fill=PALE, line=GREEN_2)
    text(slide, "ผลที่ตรวจแล้ว", 1.05, 1.9, 3.0, 0.42, size=20, bold=True, color=GREEN)
    results = [
        ("50/50", "Automated tests ผ่าน"),
        ("Pi 4", "OCR, Local LLM และเว็บเริ่มทำงานได้"),
        ("LAN", "เปิดเว็บจากอุปกรณ์ในเครือข่ายเดียวกัน"),
        ("Offline", "ไม่เรียก Cloud AI ระหว่างใช้งาน"),
    ]
    for index, (value, label) in enumerate(results):
        y = 2.56 + index * 0.82
        text(slide, value, 1.06, y, 1.48, 0.48, size=15.5, bold=True, color=GREEN_2, font=LATIN)
        text(slide, label, 2.65, y + 0.02, 3.32, 0.44, size=12.5, color=INK)

    rect(slide, 6.72, 1.58, 5.88, 4.95, fill=CREAM, line=GOLD)
    text(slide, "ข้อจำกัด", 7.05, 1.9, 2.5, 0.42, size=20, bold=True, color=RED)
    bullet_list(slide, [
        "ภาพเบลอหรือแสงสะท้อนทำให้ OCR อ่านผิด",
        "ใบเสร็จแต่ละร้านมีรูปแบบและโปรโมชั่นต่างกัน",
        "โมเดลขนาดเล็กช้ากว่า Cloud และอาจแยกข้อมูลผิด",
        "ผู้ใช้ต้องตรวจและแก้ข้อมูลก่อนบันทึก",
    ], 7.05, 2.55, 5.0, 2.95, size=13.2, color=INK, accent=RED)
    text(slide, "ระบบเป็นต้นแบบสำหรับใบเสร็จพิมพ์ชัดเจน 2–3 รูปแบบ", 7.05, 5.75, 5.0, 0.42, size=11.5, bold=True, color=GRAY)
    footer(slide)


def slide_summary(prs):
    slide = blank_slide(prs, GREEN)
    text(slide, "สรุป", 0.8, 0.72, 2.0, 0.5, size=12, bold=True, color=LIME)
    text(slide, "OCR อ่าน · LLM เข้าใจ\nผู้ใช้ตรวจ · Dashboard สรุป", 0.8, 1.35, 7.25, 1.65, size=29, bold=True, color=WHITE)
    items = [
        "ประมวลผลอยู่ภายใน Raspberry Pi",
        "ลดการส่งข้อมูลใบเสร็จไปบริการภายนอก",
        "แก้ไขผล AI ได้ก่อนบันทึกทุกครั้ง",
    ]
    for index, item in enumerate(items):
        circle_number(slide, "✓", 0.86, 3.5 + index * 0.62, diameter=0.42, fill=LIME)
        text(slide, item, 1.45, 3.5 + index * 0.62, 6.1, 0.4, size=15, color=WHITE)
    rect(slide, 8.45, 0.75, 4.05, 5.95, fill=WHITE, line=WHITE)
    text(slide, "ลำดับสาธิต", 8.82, 1.12, 3.3, 0.45, size=20, bold=True, color=GREEN)
    demo = ["อัปโหลดภาพ", "รอ OCR + AI", "ตรวจและแก้ข้อมูล", "ยืนยันบันทึก", "เปิด Dashboard"]
    for index, item in enumerate(demo):
        circle_number(slide, index + 1, 8.85, 1.9 + index * 0.78, diameter=0.44, fill=GREEN_2)
        text(slide, item, 9.46, 1.93 + index * 0.78, 2.45, 0.34, size=13.5, bold=True, color=INK)
    text(slide, "ขอบคุณครับ", 0.82, 6.42, 3.2, 0.45, size=18, bold=True, color=WHITE)
    text(slide, "Q & A", 10.05, 6.2, 1.5, 0.45, size=18, bold=True, color=GREEN, align=PP_ALIGN.CENTER)


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    slide_cover(prs)
    slide_problem(prs)
    slide_architecture(prs)
    slide_workflow(prs)
    slide_technology(prs)
    slide_llm(prs)
    slide_screens(prs)
    slide_results(prs)
    slide_summary(prs)
    prs.core_properties.title = "ใบเสร็จดี — Offline AI Receipt Analyzer"
    prs.core_properties.subject = "Project IoT presentation"
    prs.core_properties.author = "[กรอกชื่อผู้จัดทำ]"
    prs.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build()
