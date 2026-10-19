"""Generate editable Thai documents from actual screenshots and the report source."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from playwright.sync_api import sync_playwright

ROOT=Path(__file__).resolve().parents[1]
DOCS=ROOT/'docs'
EVIDENCE=DOCS/'evidence'


def format_doc(document, size=11):
    section=document.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=Cm(1.5);section.bottom_margin=Cm(1.5)
    section.left_margin=Cm(1.6);section.right_margin=Cm(1.6)
    for style_name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption','List Bullet','List Number']:
        style=document.styles[style_name]
        style.font.name='Noto Sans Thai'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans Thai')
        style.element.get_or_add_rPr().rFonts.set(qn('w:cs'),'Noto Sans Thai')
        style.font.size=Pt(size if style_name in ['Normal','List Bullet','List Number'] else 16 if 'Heading' in style_name else 20 if style_name=='Title' else 9)
        style.font.color.rgb=RGBColor.from_string('173D36')
        style.paragraph_format.space_after=Pt(5)
    document.styles['Normal'].paragraph_format.line_spacing=1.15
    footer=section.footer.paragraphs[0]
    footer.text='ใบเสร็จดี · Offline AI Receipt Analyzer   |   '
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    for run in footer.runs:run.font.size=Pt(8)


def caption(document,text):
    document.add_paragraph(text,'Caption')


def no_row_split(row):
    props=row._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))


def build_manual():
    d=Document();format_doc(d,10)
    d.add_heading('คู่มือใช้งาน ใบเสร็จดี',0)
    d.add_paragraph('สำหรับผู้ใช้งานระบบที่ติดตั้งบน Raspberry Pi แล้ว · หน้า 1/2')
    d.add_heading('1. เปิดระบบและเชื่อมต่อ',1)
    for text in [
        'เสียบไฟ Raspberry Pi และรอให้ระบบเริ่มทำงาน หากตั้งบริการอัตโนมัติแล้ว เว็บและ AI จะเริ่มเอง',
        'เชื่อมมือถือเข้ากับ Wi-Fi เดียวกับ Pi เปิดเบราว์เซอร์แล้วพิมพ์ IP ของ Pi ตามด้วย :5000 เช่น http://192.168.1.50:5000 (เปลี่ยนเป็น IP เครื่องจริง)',
        'ดูสถานะ OCR และ AI บนหน้าแรก รอให้พร้อมก่อนทดสอบ AI หาก AI ยังไม่พร้อม ระบบจะให้กรอกเองและไม่ถือว่า AI วิเคราะห์สำเร็จ'
    ]:d.add_paragraph(text,'List Number')
    d.add_heading('2. ถ่ายภาพและอัปโหลดใบเสร็จ',1)
    table=d.add_table(rows=1,cols=2);table.autofit=False
    table.columns[0].width=Cm(8);table.columns[1].width=Cm(9)
    left,right=table.rows[0].cells;no_row_split(table.rows[0])
    left.paragraphs[0].add_run().add_picture(str(EVIDENCE/'05-upload-mobile.png'),width=Cm(7.4))
    for index,text in enumerate([
        'วางใบเสร็จให้เรียบ มีแสงพอ ไม่เอียง และเห็นรายการกับยอดรวมครบ',
        'แตะ “เลือกรูปใบเสร็จของคุณ” แล้วเลือกรูป JPG หรือ PNG ขนาดไม่เกิน 10 MB / 20 ล้านพิกเซล',
        'กด “เริ่มอ่านใบเสร็จ” แล้วรอ หน้าจอจะบอกว่ากำลังวิเคราะห์ บน Pi อาจใช้เวลาหลายนาที',
        'หากออกจากหน้าเว็บ กลับไปหน้าอัปโหลดและเปิดรายการจาก “รายการที่รอตรวจ”'
    ],1):right.add_paragraph(f'{index}. {text}')
    caption(d,'ภาพที่ 1 หน้าอัปโหลดจริงบนหน้าจอมือถือ (ข้อมูลสาธิต)')
    d.add_heading('ก่อนเริ่มครั้งแรก',2)
    d.add_paragraph('ผู้ติดตั้งทำตาม README.md: รัน setup_pi.sh เปิด start_llm.sh และ start_web.sh ทดสอบก่อนติดตั้งบริการอัตโนมัติด้วย install_services.sh ระบบใช้ LAN ภายใน ไม่มีบัญชีผู้ใช้')
    d.add_page_break()
    d.add_heading('ตรวจแก้ บันทึก และดูรายจ่าย',0)
    d.add_paragraph('คู่มือใช้งาน ใบเสร็จดี · หน้า 2/2')
    d.add_heading('3. ตรวจข้อมูลก่อนยืนยัน',1)
    d.add_paragraph('เทียบภาพกับชื่อร้าน วันที่ (ปี ค.ศ.) ชื่อสินค้า จำนวน ยอดรวมบรรทัด และหมวดหมู่ แก้ไขทุกช่องที่ผิด เพิ่มหรือลบสินค้าได้ ยอดบรรทัดคือราคาทั้งบรรทัด ไม่ต้องคูณจำนวนซ้ำ')
    d.add_paragraph('ส่วนลดใส่ค่าลบ เช่น -10.00 และระบุเหตุผล ตรวจยอดรวมบนใบเสร็จ หากยอดต่าง ให้แก้หรือเลือกยืนยันส่วนต่างหลังตรวจจริง จากนั้นกด “ยืนยันและบันทึกค่าใช้จ่าย”')
    table=d.add_table(rows=1,cols=2);table.autofit=False
    for col in table.columns:col.width=Cm(8.6)
    for cell,image,label in zip(table.rows[0].cells,['06-review-mobile.png','07-dashboard-mobile.png'],['ภาพที่ 2 ตรวจยอดและยืนยัน','ภาพที่ 3 ภาพรวมรายจ่าย']):
        cell.paragraphs[0].add_run().add_picture(str(EVIDENCE/image),width=Cm(6.1))
        cell.add_paragraph(label,'Caption')
    no_row_split(table.rows[0])
    d.add_heading('4. ดู Dashboard และแก้ปัญหาเบื้องต้น',1)
    d.add_paragraph('เปิด “ภาพรวมรายจ่าย” เลือกเดือนเพื่อดูยอดรวม หมวดหมู่ และประวัติ แตะชื่อร้านเพื่อดูรายละเอียด ใบที่ยังไม่ยืนยันจะไม่ถูกนับ')
    d.add_paragraph('อ่านผิด: ถ่ายใหม่หรือกรอกเอง • เข้าเว็บไม่ได้: ตรวจ Wi-Fi, IP และไฟ Pi • AI ไม่พร้อม: ให้ผู้ติดตั้งตรวจบริการ • ก่อนถอดไฟให้สั่งปิดเครื่องอย่างถูกต้อง ข้อมูลเก็บบน Pi ควรสำรองตาม README')
    d.save(DOCS/'คู่มือใช้งาน.docx')


def build_report():
    d=Document();format_doc(d)
    lines=(DOCS/'REPORT.md').read_text().splitlines()
    table_lines=[];in_code=False
    def flush_table():
        if not table_lines:return
        rows=[[c.strip() for c in line.strip('|').split('|')] for line in table_lines if not re.match(r'^\|[\s:|\-]+\|$',line)]
        table=d.add_table(rows=0,cols=len(rows[0]));table.style='Light Shading Accent 1'
        for data in rows:
            cells=table.add_row().cells
            for cell,text in zip(cells,data):cell.text=text
        table_lines.clear()
    for line in lines:
        if line.startswith('|'):
            table_lines.append(line);continue
        flush_table()
        if line.startswith('```'):in_code=not in_code;continue
        if not line.strip():continue
        if in_code:
            p=d.add_paragraph();r=p.add_run(line);r.font.name='DejaVu Sans Mono';r.font.size=Pt(9)
        elif line.startswith('# '):d.add_heading(line[2:],0)
        elif line.startswith('## '):d.add_heading(line[3:],1)
        elif line.startswith('- '):d.add_paragraph(line[2:],'List Bullet')
        else:d.add_paragraph(line)
        if line.startswith('ภาพประกอบหลัก:'):
            d.add_picture(str(DOCS/'architecture.png'),width=Cm(17.5));caption(d,'ภาพที่ 1 สถาปัตยกรรมและลำดับการทำงาน')
        if line.startswith('ผล cloud ตรวจ'):
            for image,label in [('01-upload.png','หน้าอัปโหลด'),('02-review.png','หน้าตรวจแก้'),('03-dashboard.png','Dashboard')]:
                d.add_picture(str(EVIDENCE/image),width=Cm(16.5));caption(d,f'{label} — ข้อมูลสาธิตจาก browser smoke test')
    flush_table()
    d.save(DOCS/'รายงานโครงงาน-ฉบับร่าง.docx')


if __name__=='__main__':
    with sync_playwright() as p:
        browser=p.chromium.launch(executable_path='/usr/bin/chromium',headless=True,args=['--no-sandbox'])
        page=browser.new_page(viewport={'width':1400,'height':740})
        page.set_content('<html><body style="margin:0">'+(DOCS/'architecture.svg').read_text()+'</body></html>');page.evaluate('document.fonts.ready')
        page.screenshot(path=str(DOCS/'architecture.png'))
        browser.close()
    build_manual();build_report()
    print('Generated editable manual and draft report DOCX files.')
